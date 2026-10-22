"""Monta a pasta de entrega para o orientador (16/10/2026): um .docx com os itens que ele marcou
em VERMELHO no texto (reunião de 02/10/2026), na ordem e numeração do texto DELE, com as tabelas
como tabelas do Word (editáveis) e as figuras embutidas; e as figuras e tabelas soltas na pasta,
para ele colocar onde precisar.

Uso (raiz do projeto, .venv ativo):
    python src/gerar_entrega_orientador.py [pasta_execucao_completa]

Saída: <pasta_execucao>/entrega_orientador_<AAAA-MM-DD>/
    entrega_orientador.docx          documento com os entregáveis
    figuras/                         PNG (versão colorida e em preto e branco) de cada figura citada
    tabelas/                         .tex e .png de cada tabela citada
    variantes_de_prompt.txt          texto real dos quatro enunciados (padrão, geométrico, cot, tabular)
    csv/                             CSVs de onde saem os números
A seleção de itens está em ITENS (um por vermelho do texto); os textos de legenda são de uma linha,
sem análise: o texto é do orientador.
"""
import glob
import os
import re
import shutil
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

import gerar_tabelas as gt  # noqa: E402
from tabelas import latex_para_texto  # noqa: E402

ALIAS = {"openai/gpt-4o-mini": "gpt4mini", "google/gemini-2.5-flash-lite": "flashlite"}
NOME_MODELO = {"gpt4mini": "gpt-4o-mini", "flashlite": "gemini-2.5-flash-lite"}
MODELOS = ["gpt4mini", "flashlite"]     # gpt primeiro (narrativa principal), Gemini como contraste

# Cada item: título com a numeração DO TEXTO DELE, a frase em vermelho, tabelas (prefixo sem alias),
# figuras (prefixo; seeds e painéis resolvidos por glob) e uma observação de 1 a 3 linhas.
ITENS = [
    {
        "id": "R1", "titulo": "Tabela 2 (V.1): fidelidade da métrica no Problema A, zero-shot e few-shot",
        "vermelho": "Tabela 2: Fidelidade da métrica no Problema A, com 3 sementes e zero-shot. Falta fazer experimentos com few-shot e gerar uma nova tabela.",
        "tabelas": ["bloco1_table_01_A_passos_2_a_5", "bloco1_table_02_A_passo2_vs_passo3", "bloco1_table_05_A_passo4_perceptron_vs_nnls"],
        "figuras": [],
        "obs": ["Few-shot do Passo 3: k ∈ {4, 10, 20, 40} exemplos rotulados pelo rótulo real (k/2 por classe, 3 sorteios por k); o LLM classifica os 150 − k restantes.",
                "Em zero-shot a correspondência nome↔classe é arbitrária (o gpt-4o-mini rotulou 109/41): a acurácia literal do LLM fica abaixo de 50 %; o valor invariante à permutação é max(a, 1 − a).",
                "† = o Perceptron não encontrou Ŵ que separe os rótulos do LLM (γ = 0 ou margem bruta mínima negativa); regra de 15/10/2026."],
    },
    {
        "id": "R2", "titulo": "Figura pedida em V.1: amostras do Problema A rotuladas pelo rótulo real, pelo LLM e pela métrica",
        "vermelho": "Seria interessante plotar as amostras das gaussianas rotuladas pelo LLM e pelo ground truth.",
        "tabelas": [],
        "figuras": ["bloco1_graphic_02_A_passos_2_a_4_simbolos_seed42", "bloco1_graphic_11_B_passos_2_3_simbolos_seed42",
                    "bloco1_graphic_11_C_passos_2_3_simbolos_seed42", "bloco1_graphic_13_D_passos_2_a_4_simbolos_seed42"],
        "obs": ["Símbolos + (classe 0) e − (classe 1), centróides como o, exemplos do few-shot circulados; versão em preto e branco. Semente 42 aqui; as sementes 123 e 7 estão na pasta figuras/.",
                "A mesma figura para B e C (com os centróides geradores e a rotação em relação ao A) e para a meia-lua (Problema D)."],
    },
    {
        "id": "R3", "titulo": "Figura 1 (V.1): concordância dos estimadores, fidelidade no A e consistência em B e C",
        "vermelho": "Figura 1: Bloco 1: comparação da concordância dos estimadores, fidelidade da métrica no Problema A e consistência do LLM nas transferências, Problema B e Problema C.",
        "tabelas": [],
        "figuras": ["bloco1_graphic_06_A_passo4_perceptron_vs_nnls", "bloco1_graphic_10_BC_passo3_consistencia"],
        "obs": ["Perceptron × NNLS pareados sobre os mesmos rótulos do LLM (0 chamadas); consistência, Cohen's κ e F1-Macro em B e C com a métrica zero-shot do A e os centróides locais."],
    },
    {
        "id": "R4", "titulo": "Tabela 3 (V.2): teste-oráculo, Problemas 1, 2 e 3",
        "vermelho": "Tabela 3: Teste-oráculo: recuperação da métrica e concordância entre estimadores. Qual a diferença entre problema A e B? As métricas foram utilizadas para rotular o problema A? Renomear problemas Expert aniso. Nome muito esquisito. Sugestão: Problema 1, Problema 2, Problema 3 (e 4) com as métricas na frente.",
        "tabelas": ["bloco1_table_17_oraculo_passos_1_a_3"],
        "figuras": [],
        "obs": ["Resposta: as três métricas, [1; 1] (Problema 1), [0,3; 1,5] (Problema 2) e [1,5; 0,3] (Problema 3), rotulam o MESMO conjunto, as gaussianas do Problema E (centróides (−1,5; 1) e (1,5; −1), 150 pontos, σ = 1,3), pelo centróide mais próximo.",
                "O Problema A não é usado: com os centróides em x2 = 0 (até 14/10/2026) a diferença de distâncias não dependia de w2 e a métrica não era identificável. A primeira linha da tabela antiga foi retirada, como pedido."],
    },
    {
        "id": "R5", "titulo": "Figura 2 (V.2): validação dos estimadores, em quatro painéis",
        "vermelho": "Figura 2: Validação dos Estimadores (quebrar; está pequena).",
        "tabelas": [],
        "figuras": ["bloco1_graphic_19_oraculo_passos_1_a_3_w_scatter", "bloco1_graphic_19_oraculo_passos_1_a_3_ratio",
                    "bloco1_graphic_19_oraculo_passos_1_a_3_cosseno", "bloco1_graphic_19_oraculo_passos_1_a_3_acuracia"],
        "obs": ["Os quatro painéis separados (W recuperado, razão w1/w2 em escala log, cosseno, acurácia), média e faixa mín–máx das 3 sementes."],
    },
    {
        "id": "R6", "titulo": "Tabela 4 (V.3): transferência para B e C, consistência e κ, zero-shot e few-shot",
        "vermelho": "Tabela 4: Transferência: consistência e κ (Perceptron), IC bootstrap 95%.",
        "tabelas": ["bloco1_table_08_BC_passo3_zero_vs_eleita", "bloco1_table_09_BC_passo4_exemplos_da_metrica", "bloco1_table_07_BC_passo3_por_semente"],
        "figuras": [],
        "obs": ["Sem IC bootstrap (retirado a pedido de 13/10/2026): média e faixa mín–máx entre as sementes.",
                "Few-shot em B e C com k ∈ {4, 10, 20, 40} exemplos rotulados pela métrica do A (não só 4)."],
    },
    {
        "id": "R7", "titulo": "V.4 e V.7: teto da métrica diagonal na meia-lua (substitui os \"100 % e 95,3 %\")",
        "vermelho": "No teste-oráculo no problema Meia-Lua os estimadores recuperam a métrica de forma correta mesmo em geometria curva atingindo uma alta acurácia, Perceptron Estruturado 100% e NNLS 95,3%.",
        "tabelas": ["bloco1_table_18_D_passo5_teto"],
        "figuras": ["bloco1_graphic_21_D_passo5_teto"],
        "obs": ["Não há 100 %: a métrica diagonal ajustada diretamente ao rótulo real da meia-lua chega a 86,9 % (R²), 87,1 % (R³) e 87,3 % (R⁴) com o Perceptron e 86,4 a 86,7 % com o NNLS. A frase de V.7 (86,7 %) é a correta."],
    },
    {
        "id": "R8", "titulo": "Tabela 5 (V.4): fidelidade em R², R³ e R⁴ no Problema A",
        "vermelho": "Tabela 5: Fidelidade média do Perceptron Estruturado em R3 e R4 para os problemas A, B e C.",
        "tabelas": ["bloco1_table_04_A_passo4_R2_R3_R4"],
        "figuras": ["bloco1_graphic_08_A_passo4_R2_R3_R4"],
        "obs": ["Só o Problema A: B e C são lineares e recebem apenas o vetor w do A, logo R³/R⁴ não se aplica a eles (decisão de 13/10/2026)."],
    },
    {
        "id": "R9", "titulo": "V.4: resultados e tabelas do Problema D (meia-lua), mesmos passos do A",
        "vermelho": "Falta resultados e tabela relativa ao problema D.",
        "tabelas": ["bloco1_table_10_D_passos_2_a_5", "bloco1_table_11_D_passo2_vs_passo3", "bloco1_table_12_AD_por_passo"],
        "figuras": ["bloco1_graphic_04_AD_passos_2_a_5_curva", "bloco1_graphic_14_D_passo6_svm_seed42"],
        "obs": ["Zero-shot, few-shot do rótulo real (k ∈ {4, 10, 20, 40}) e teto, em R², R³ e R⁴, sobre os 150 pontos.",
                "O gpt-4o-mini colapsa em zero-shot na meia-lua nas 3 sementes (uma classe para todos os pontos); com exemplos deixa de colapsar. O Gemini não colapsa."],
    },
    {
        "id": "R10", "titulo": "Tabela 6 (V.5): nomes de classe, nomes de atributo e variantes de prompt",
        "vermelho": "Tabela 6 - Efeito da nomenclatura e variante de prompt sobre a fidelidade. (\"corrigir nomes peso/altura\")",
        "tabelas": ["bloco1_table_14_vies_nomes_classe", "bloco1_table_15_vies_nomes_atributo", "bloco1_table_16_vies_variantes_prompt"],
        "figuras": ["bloco1_graphic_16_vies_nomes_classe", "bloco1_graphic_15_vies_ordem_classes", "bloco1_graphic_17_vies_nomes_atributo"],
        "obs": ["Nomes de atributo corrigidos (altura/peso é a variante semântica do Problema A; peso/altura é do Problema G). Os valores do texto (99,3 %, 82,7 → 68,0 %, 84,7 %, 68 %, 64,7 %) eram da execução de julho e mudaram."],
    },
    {
        "id": "R11", "titulo": "Figura 3 (V.5): variantes de prompt, com a explicação de cada variante",
        "vermelho": "Figura 3 - Fidelidade e consistência da métrica em relação a variantes de prompt. \"precisa explicar como é feita cada variante\"",
        "tabelas": [],
        "figuras": ["bloco1_graphic_18_vies_variantes_prompt"],
        "obs": ["As quatro variantes estão descritas na seção 3 deste documento, com o texto real de cada enunciado (também em variantes_de_prompt.txt)."],
    },
    {
        "id": "R12", "titulo": "Tabela 7 (V.6): o valor repetido em zero-shot",
        "vermelho": "A acurácia média, sobre as quatro formas de seleção de exemplos é de 61% \"estranho o valor se repetir\".",
        "tabelas": ["bloco2_table_04_E_passo2_vs_passo3"],
        "figuras": [],
        "obs": ["Não é erro: em zero-shot não há exemplos, logo não há estratégia de seleção; é UMA coleta por (modelo, semente), repetida nas quatro linhas. Valores atuais (E1): 61,3 % (gpt-4o-mini) e 54,4 % (Gemini)."],
    },
    {
        "id": "R13", "titulo": "Tabela 7 e Figura 4 (V.6): aprendizado por estratégia de seleção, E1, E2 e E3",
        "vermelho": "A Tabela 7 apresenta a variação do aprendizado para cada forma de seleção considerando zero-shot e few-shot para n ∈ {4, 10, 20, 40}. Figura 4: Bloco 2 - Curva de aprendizado por estratégia de seleção de exemplos. (renomear \"aniso_x2\")",
        "tabelas": ["bloco2_table_01_E1_passo3_estrategia_x_k", "bloco2_table_02_E2_passo3_estrategia_x_k", "bloco2_table_03_E3_passo3_estrategia_x_k", "bloco2_table_05_E_passo3_H4"],
        "figuras": ["bloco2_graphic_03_E1_passo3_curva", "bloco2_graphic_04_E1_passo3_estrategias"],
        "obs": ["Nomes atuais: E1 = [0,3; 1,5] (x2 cinco vezes mais relevante), E2 = [1,5; 0,3], E3 = [1; 1].",
                "H4 (difícil > fácil) refutada nos dois modelos; E1, média dos k > 0: gpt-4o-mini aleatória 89,9 > fácil 88,1 > mista 86,7 > difícil 77,7 %; Gemini fácil 93,2 > aleatória 92,9 > mista 91,1 > difícil 72,4 %."],
    },
    {
        "id": "R14", "titulo": "Tabela 8 (V.7): LLM com 40 exemplos versus k-NN, Regressão Logística e SVM",
        "vermelho": "Tabela 8 - Bloco 2 - LLM (40-shot) versus baselines.",
        "tabelas": ["bloco2_table_06_E_passo6_baselines"],
        "figuras": ["bloco2_graphic_05_E1_passo6_baselines"],
        "obs": ["Classificadores treinados nos MESMOS exemplos mostrados ao LLM (estratégia mista, 3 sorteios), avaliados nos mesmos pontos; E1, E2 e E3 (antigas variantes P3 e P4)."],
    },
    {
        "id": "R15", "titulo": "Figura 6 (V.7): meia-lua, SVM RBF versus zero-shot do LLM",
        "vermelho": "Figura 6 - Problema F (Meia-Lua, seed 42, gpt-4o-mini): superfície de separação do SVM RBF treinado no ground truth versus zero-shot do LLM.",
        "tabelas": ["bloco1_table_19_D_passo6_svm", "bloco2_table_10_F_passo3_k"],
        "figuras": ["bloco1_graphic_14_D_passo6_svm_seed42"],
        "obs": ["Mudou: nesta execução o gpt-4o-mini colapsa em zero-shot nas 3 sementes da meia-lua (não só em 123 e 7), e o painel dele registra o colapso; a figura com fronteira é a do Gemini. A tabela traz o Problema F (split 70/30) com k ∈ {0, 4, 10, 20, 40}: os números de V.7 (53,3 % → 68 %) precisam ser substituídos."],
    },
    {
        "id": "R16", "titulo": "Figura 7 (V.7): diluição (esquerda) e viés de ordem dos exemplos (direita)",
        "vermelho": "Figura 7 - Esquerda: experimento de diluição (acrescentar exemplos fáceis a três exemplos difíceis melhora de forma limitada a acurácia). Direita: viés de ordem (as quatro ordenações diferem por poucos pontos percentuais).",
        "tabelas": ["bloco2_table_07_E1_passo4_diluicao", "bloco2_table_08_E1_passo5_ordem"],
        "figuras": ["bloco2_graphic_06_E1_passo4_diluicao", "bloco2_graphic_07_E1_passo5_ordem"],
        "obs": ["São QUATRO exemplos difíceis fixos (não três) mais 0, 2, 4, 10, 16 e 20 fáceis. Ordenações: classe 0 primeiro, classe 1 primeiro, embaralhada e alternada, com k ∈ {4, 10, 20}."],
    },
]

NUMEROS = [
    ("V (abertura) e IV.2", "129.868 chamadas; limite de dez consultas em paralelo; 0 respostas não interpretadas; gpt-4o-mini muda 4,3 % das respostas repetidas",
     "222.648 chamadas; 50 em paralelo; 5 respostas sem rótulo interpretável (0,002 %, todas do Gemini na variante cot, por tempo limite de 90 s); 1,19 % no gpt-4o-mini e 0 % no Gemini"),
    ("IV.1, Tabela 1", "D, E e F com \"? amostras\"; A com centróides (−2, 0) e (2, 0); B e C com rotação ±1,5",
     "D 150; E 150 (centróides (−1,5; 1) e (1,5; −1), σ = 1,3); F 150 (105 de treino e 45 de teste); G 100 (50/50), 70/30. Desde 15/10/2026 o A tem centróides (−2; −0,5) e (2; 0,5), σ = 1,2; B (−2; 1,5)/(2; −1,5) e C (−2; −1,5)/(2; 1,5) continuam, logo a rotação em relação ao A é 2,0 em B (horária) e 1,0 em C (anti-horária)"),
    ("IV.2", "few-shot n ∈ {4, 10} (especialista) e {4, 10, 20, 40} (aprendiz)", "{4, 10, 20, 40} nos dois papéis (Problemas A, D e G no Passo 3; B e C no Passo 4; E, F e G no aprendizado em contexto)"),
    ("IV.7", "intervalos de confiança de 95 % via bootstrap com 10⁴ amostragens", "retirado (pedido de 13/10/2026): média e faixa mín–máx entre as sementes, diferenças pareadas por semente e Cohen's d"),
    ("IV.7", "\"Este mesmo protocolo de divisão [70/30] é aplicado aos problemas E e F\"", "F tem a divisão 70/30; E não: o LLM classifica os 150 − k pontos restantes. Decidir se a frase muda ou se o E passa a ter divisão"),
    ("V.1", "fidelidade de 81 % a 86 %; o LLM acerta de 25 % a 31 % do rótulo real; cosseno médio 0,986", "gpt-4o-mini, zero-shot: fidelidade 86,7 % (85,3 a 88,0); acurácia literal 26,9 % (24,7 a 28,0; invariante 73 %); Tabela 2/5 desta entrega"),
    ("V.3", "κ = 0,33 em B (razoável) e 0,75 em C (substancial), zero-shot", "κ_B = 0,320 (gpt-4o-mini) e 0,258 (Gemini); κ_C = 0,725 e 0,852; com k ≥ 10 exemplos, κ_B sobe para 0,85 a 0,89"),
    ("V.4", "média sobe de 84,9 % (3 atributos) para 86,2 % (4); oráculo na meia-lua 100 % e 95,3 %", "ver Tabela 5 desta entrega; teto da meia-lua 86,9 / 87,1 / 87,3 % (R², R³, R⁴)"),
    ("V.6", "Aleatória 90,3 > Mista 88,0 > Fácil 81,1 > Difícil 75,1; Gemini Fácil 93,5 vs Difícil 78,1", "gpt-4o-mini (E1): aleatória 89,9 > fácil 88,1 > mista 86,7 > difícil 77,7; Gemini: fácil 93,2 > aleatória 92,9 > mista 91,1 > difícil 72,4"),
    ("V.7", "gpt-4o-mini colapsa em 2 das 3 sementes (123 e 7); na semente 42 parte de 53,3 % e satura em ≈ 68 %", "colapsa nas 3 sementes em zero-shot (Problemas D e F); com exemplos classifica: ver Tabela do Problema F nesta entrega"),
    ("Fundamentação, Algoritmo 1", "pseudocódigo validado na reunião de 02/10", "três ajustes em 15/10/2026, sem mudar o critério de parada (margem relaxada com α): w e α divididos por ‖w‖ ao fim de cada época; melhor esforço = época com menos pontos mal classificados; margem bruta mínima registrada (marcador †). Documentos: algoritmo_codigo_executado_pseudocodigo.docx e algoritmo_perceptron_animado.html"),
]

PERGUNTAS = [
    "Problema E: o texto diz que a divisão 70/30 vale para E e F; no sistema só o F tem divisão (o E classifica os 150 − k restantes). A frase muda ou o E passa a ter divisão?",
    "Grade de exemplos: o sistema usa k ∈ {4, 10, 20, 40} em todos os passos com exemplos; o texto diz {4, 10} para o especialista. Manter a grade cheia nas tabelas?",
    "Tabela 5: fica só com o Problema A (B e C não recebem R³/R⁴).",
    "Meia-lua: o gpt-4o-mini colapsa em zero-shot nas três sementes; o Gemini não. Os parágrafos de V.4 e V.7 que citam \"100 %\" e \"semente 42\" precisam ser reescritos.",
    "Problema A: os centróides passaram a (−2; −0,5) e (2; 0,5) em 15/10/2026, porque com x2 = 0 nos dois centros o peso w2 não era identificável e o κ em B oscilava entre execuções. Todos os números do Bloco 1 são desta configuração.",
    "Perceptron: os três ajustes de 15/10/2026 (normalização por época, melhor esforço por erros, margem bruta mínima) estão no pseudocódigo e na página animada enviados; o critério de parada do texto não mudou.",
    "Feedback imediato: continua adiado, como combinado na reunião.",
    "Bloco 3 (peso × altura): as tabelas e figuras já existem (bloco3_table_01 a 06, bloco3_graphic_01 a 03, bloco23_graphic_01 a 04) e podem ser enviadas quando for escrever essa parte.",
]


def _execucao_mais_recente() -> Path:
    pastas = sorted(glob.glob(str(RAIZ / "resultados" / "execucao_*_completa")))
    if not pastas:
        raise SystemExit("Nenhuma pasta resultados/execucao_*_completa.")
    return Path(pastas[-1])


def _par(doc, texto, size=10, bold=False, italic=False, cor=None, align=None, antes=0, depois=4):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if cor:
        r.font.color.rgb = RGBColor(*cor)
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(antes); p.paragraph_format.space_after = Pt(depois)
    return p


def _tabela_word(doc, t):
    """Tabela do catálogo (objeto Tabela) → tabela do Word, editável."""
    heads = [latex_para_texto(h) for h in t.headers]
    ncols = len(heads)
    _par(doc, latex_para_texto(t.caption), size=9.5, bold=True, antes=6, depois=2)
    tab = doc.add_table(rows=0, cols=ncols)
    tab.style = "Table Grid"
    if t.grupos:
        cells = tab.add_row().cells
        for titulo, c0, c1 in t.grupos:
            a, b = cells[c0 - 1], cells[c1 - 1]
            m = a.merge(b) if c1 > c0 else a
            m.text = ""
            run = m.paragraphs[0].add_run(latex_para_texto(titulo)); run.bold = True; run.font.size = Pt(8)
            m.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells = tab.add_row().cells
    for c, h in zip(cells, heads):
        c.text = ""
        run = c.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(8)
    for i, row in enumerate(t.rows):
        cells = tab.add_row().cells
        for j, (c, v) in enumerate(zip(cells, row)):
            c.text = ""
            run = c.paragraphs[0].add_run(latex_para_texto(str(v)))
            run.font.size = Pt(8)
            if i in t.destaque:
                run.bold = True
            if j > 0:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if t.rodape:
        _par(doc, latex_para_texto(t.rodape), size=7.5, italic=True, antes=2, depois=8)


def _figura(doc, caminho: Path, legenda: str, largura_cm: float = 24.0):
    doc.add_picture(str(caminho), width=Cm(largura_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _par(doc, legenda, size=8, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, depois=8)


def _variantes_de_prompt():
    import dissertacao_mestrado as dm
    saida = []
    descr = {
        "default": "Padrão: enunciado direto, o ponto com seus atributos e a pergunta pela classe.",
        "geometric": "Geométrico: o enunciado descreve explicitamente o ponto como posição num plano cartesiano (contexto espacial).",
        "cot": "Chain-of-thought (cot): pede ao modelo raciocinar passo a passo antes de dar o rótulo (o rótulo é extraído do fim da resposta).",
        "tabular": "Tabular: os atributos do ponto (e dos exemplos, quando há) são apresentados numa tabela em markdown.",
    }
    for v in ("default", "geometric", "cot", "tabular"):
        zero = dm._montar_prompt(1.25, -0.5, "A", "B", None, "x1", "x2", None, v)
        few = dm._montar_prompt(1.25, -0.5, "A", "B", [(-2.1, 0.3, "A"), (1.9, -0.2, "B")], "x1", "x2", None, v)
        sysmsg = dm.PROMPT_VARIANTS.get(v, {}).get("system_message", "")
        saida.append((v, descr[v], sysmsg, zero, few))
    return saida


def main():
    execucao = Path(sys.argv[1]) if len(sys.argv) > 1 else _execucao_mais_recente()
    hoje = date.today()
    out = execucao / f"entrega_orientador_{hoje.isoformat()}"
    for sub in ("figuras", "tabelas", "csv"):
        (out / sub).mkdir(parents=True, exist_ok=True)

    D = gt.carregar_tudo(execucao); D["_pasta"] = execucao
    tabs = gt.catalogo(D, None)
    por_nome = {}
    for t in tabs:
        por_nome[t.nome_png().replace("_gray.png", "")] = t

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10)

    _par(doc, "Entrega: tabelas e figuras marcadas em vermelho no texto (reunião de 02/10/2026)", size=16, bold=True)
    _par(doc, f"George Gomes, {hoje.strftime('%d/%m/%Y')}. Fonte: execução completa de 15/10/2026 à noite ({execucao.name}): "
              "gpt-4o-mini e gemini-2.5-flash-lite, T = 0, sementes {42, 123, 7}, 3 repetições onde há sorteio de exemplos, 222.648 chamadas, "
              "Perceptron Estruturado com os ajustes de 15/10/2026.", size=9.5)
    _par(doc, "Como ler: a seção 1 segue a numeração do SEU texto (Tabela 2, Figura 1, …); em cada item vem a frase em vermelho, "
              "a tabela ou figura e uma legenda de uma linha, sem análise. As tabelas são tabelas do Word (editáveis); as figuras "
              "estão também soltas na pasta figuras/ (versão colorida e em preto e branco) e as tabelas em tabelas/ (.tex e .png). "
              "A seção 2 lista os números do texto que mudaram; a seção 3 explica as variantes de prompt; a seção 4 reúne as perguntas.", size=9.5)

    # índice de correspondência
    _par(doc, "Índice de correspondência", size=12, bold=True, antes=8)
    idx = doc.add_table(rows=1, cols=4); idx.style = "Table Grid"
    for c, h in zip(idx.rows[0].cells, ("No seu texto", "Item", "O que está nesta entrega", "Status")):
        c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5)
    status = {"R15": "mudou", "R11": "feito (explicação na seção 3)"}
    for it in ITENS:
        cells = idx.add_row().cells
        vals = (it["titulo"].split(":")[0], it["id"], ", ".join(it["tabelas"] + it["figuras"]) or "—", status.get(it["id"], "feito"))
        for c, v in zip(cells, vals):
            c.text = ""; r = c.paragraphs[0].add_run(v); r.font.size = Pt(8)
    for ref, item, st in (("Tabela 1 (IV.1)", "números na seção 2", "números prontos"), ("IV.7, divisão 70/30 no E", "pergunta na seção 4", "decisão sua"),
                          ("IV.7, feedback imediato", "—", "adiado, como combinado")):
        cells = idx.add_row().cells
        for c, v in zip(cells, (ref, "M", item, st)):
            c.text = ""; r = c.paragraphs[0].add_run(v); r.font.size = Pt(8)

    # seção 1
    doc.add_page_break()
    _par(doc, "1. Itens em vermelho, na ordem do texto", size=14, bold=True)
    copiados = set()
    for it in ITENS:
        _par(doc, f"{it['id']}. {it['titulo']}", size=12, bold=True, antes=10)
        _par(doc, "No texto: " + it["vermelho"], size=9, italic=True, cor=(200, 0, 0))
        for pref in it["tabelas"]:
            achou = False
            for alias in MODELOS + [None]:
                chave = f"{pref}__{alias}" if alias else pref
                t = por_nome.get(chave)
                if t is None:
                    continue
                achou = True
                _tabela_word(doc, t)
                for ext in ("tex", "png"):
                    src = (execucao / "tabelas" / t.nome_tex()) if ext == "tex" else (execucao / t.nome_png())
                    if src.exists():
                        shutil.copy(src, out / "tabelas" / src.name); copiados.add(src.name)
            if not achou:
                _par(doc, f"(tabela {pref} não encontrada nesta execução)", size=8, italic=True)
        for pref in it["figuras"]:
            for alias in MODELOS + [None]:
                suf = f"__{alias}" if alias else ""
                for cor in ("color", "gray"):
                    for f in sorted(execucao.glob(f"{pref}{suf}_{cor}.png")):
                        shutil.copy(f, out / "figuras" / f.name); copiados.add(f.name)
                cand = sorted(execucao.glob(f"{pref}{suf}_color.png"))
                if pref.startswith(("bloco1_graphic_02", "bloco1_graphic_11", "bloco1_graphic_13")):
                    cand = sorted(execucao.glob(f"{pref}{suf}_gray.png"))     # figuras de símbolos: versão P&B
                for f in cand:
                    nome = f.name
                    legenda = nome.replace("_color.png", "").replace("_gray.png", "")
                    modelo = NOME_MODELO.get(alias, "") if alias else "sem LLM"
                    _figura(doc, f, f"{legenda} ({modelo}). Arquivo em figuras/.")
                if alias is None:
                    break
            # as demais sementes (figura combinada, sem painéis) vão só para a pasta
            base = re.sub(r"_seed42$", "", pref)
            rx = re.compile(r"^" + re.escape(base) + r"(_seed\d+)?(__(gpt4mini|flashlite))?_(color|gray)\.png$")
            for f in execucao.glob(f"{base}*.png"):
                if rx.match(f.name) and f.name not in copiados:
                    shutil.copy(f, out / "figuras" / f.name); copiados.add(f.name)
        for o in it["obs"]:
            _par(doc, "• " + o, size=9, depois=2)

    # seção 2
    doc.add_page_break()
    _par(doc, "2. Números do texto que mudaram", size=14, bold=True)
    _par(doc, "Fonte: execução completa de 15/10/2026 à noite. Coluna \"deve ficar\" com os valores do gpt-4o-mini (narrativa principal) e do Gemini quando diferem.", size=9)
    tb = doc.add_table(rows=1, cols=3); tb.style = "Table Grid"
    for c, h in zip(tb.rows[0].cells, ("Onde no texto", "Está", "Deve ficar")):
        c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5)
    for onde, esta, deve in NUMEROS:
        cells = tb.add_row().cells
        for c, v in zip(cells, (onde, esta, deve)):
            c.text = ""; r = c.paragraphs[0].add_run(v); r.font.size = Pt(8.5)

    # seção 3
    doc.add_page_break()
    _par(doc, "3. As quatro variantes de prompt (pedido da reunião e do e-mail de 05/10)", size=14, bold=True)
    _par(doc, "Todas as variantes recebem o mesmo ponto e os mesmos exemplos; muda só a forma do enunciado. O exemplo abaixo usa o ponto (1,25; −0,5), classes A e B e, no few-shot, dois exemplos.", size=9)
    linhas_txt = []
    for v, d, sysmsg, zero, few in _variantes_de_prompt():
        _par(doc, d, size=10, bold=True, antes=8)
        if sysmsg:
            _par(doc, "Mensagem de sistema: " + sysmsg, size=8.5, italic=True)
        _par(doc, "Zero-shot:", size=8.5, bold=True, depois=0)
        p = _par(doc, zero, size=8, depois=4); p.runs[0].font.name = "Consolas"
        _par(doc, "Few-shot (2 exemplos):", size=8.5, bold=True, depois=0)
        p = _par(doc, few, size=8, depois=4); p.runs[0].font.name = "Consolas"
        linhas_txt += [f"===== {v}: {d}", f"[sistema] {sysmsg}", "--- zero-shot ---", zero, "--- few-shot (2 exemplos) ---", few, ""]
    (out / "variantes_de_prompt.txt").write_text("\n".join(linhas_txt), encoding="utf-8")

    # seção 4
    doc.add_page_break()
    _par(doc, "4. Perguntas e avisos", size=14, bold=True)
    for i, q in enumerate(PERGUNTAS, 1):
        _par(doc, f"{i}. {q}", size=9.5, depois=4)

    # CSVs de origem
    for f in execucao.glob("*.csv"):
        shutil.copy(f, out / "csv" / f.name)
    shutil.copy(execucao / "tabelas" / "INDICE.md", out / "tabelas" / "INDICE.md")

    destino = out / "entrega_orientador.docx"
    doc.save(destino)
    n_fig = len(list((out / "figuras").glob("*.png"))); n_tab = len(list((out / "tabelas").glob("*.tex")))
    print(f"Pasta: {out}\n  {destino.name} ({destino.stat().st_size / 1024:.0f} KB); figuras/: {n_fig} PNG; tabelas/: {n_tab} .tex; csv/: {len(list((out / 'csv').glob('*.csv')))} CSV")


if __name__ == "__main__":
    main()
