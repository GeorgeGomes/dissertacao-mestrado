"""Gera `algoritmo_codigo_executado_pseudocodigo.docx` (raiz do projeto) a partir de
`pseudocodigo.py`: pseudocódigo linha a linha do método `fit`, valores usados nas
execuções, observações e a tabela de correspondência com as linhas de
`src/relaxed_perceptron.py` (o código de cada linha é lido do arquivo atual).

Uso (raiz do projeto, .venv ativo):
    python src/animacao_perceptron/gerar_docx.py
"""
import os
import sys
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.abspath(os.path.join(AQUI, "..", ".."))
sys.path.insert(0, AQUI)

import pseudocodigo as pc  # noqa: E402

SAIDA = os.path.join(RAIZ, "algoritmo_codigo_executado_pseudocodigo.docx")


def _par(doc, texto, bold=False, size=11, mono=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = bold
    r.font.size = Pt(size)
    if mono:
        r.font.name = "Consolas"
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(4)
    return p


def _tabela(doc, linhas, larguras_cm, mono_cols=(), header=None):
    t = doc.add_table(rows=0, cols=len(larguras_cm))
    t.style = "Table Grid"
    if header:
        cells = t.add_row().cells
        for c, h in zip(cells, header):
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(10)
    for linha in linhas:
        cells = t.add_row().cells
        for j, (c, v) in enumerate(zip(cells, linha)):
            c.text = ""
            partes = str(v).split("\n")
            for k, parte in enumerate(partes):
                p = c.paragraphs[0] if k == 0 else c.add_paragraph()
                r = p.add_run(parte)
                r.font.size = Pt(9.5)
                if j in mono_cols:
                    r.font.name = "Consolas"
    for row in t.rows:
        for c, w in zip(row.cells, larguras_cm):
            c.width = Cm(w)
    return t


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    _par(doc, "Algoritmo do código executado (relaxed_perceptron.py)", bold=True, size=14)
    _par(doc, "Este documento traz o pseudocódigo do método fit da classe RelaxedPerceptron, no arquivo "
              "src/relaxed_perceptron.py. É esse código que gera os resultados do trabalho; a versão aqui descrita é a "
              "de 15/10/2026 (normalização de w e α por época, melhor esforço por pontos mal classificados e margem "
              "bruta mínima), usada na execução completa de 15/10/2026 à noite. "
              f"Documento gerado em {date.today().strftime('%d/%m/%Y')} por src/animacao_perceptron/gerar_docx.py.")
    _par(doc, "O pseudocódigo segue o código linha a linha, sem resumir nenhum passo, e usa a mesma notação da versão "
              "revisada do Algoritmo 1. A seção 4 indica, para cada linha, o trecho correspondente do código.")

    _par(doc, "1. Pseudocódigo", bold=True, size=12)
    _par(doc, "d_w(v,c) = Σ_j w_j.(v_j – c_j)², sem raiz quadrada. d é o número de atributos: d = 2 nos problemas em R², "
              "d = 3 e d = 4 nos experimentos com R³ e R⁴.")
    _par(doc, "Os centróides c(0) e c(1) são a média dos pontos de cada classe segundo os rótulos y(i) recebidos. No "
              "Passo 4 dos Problemas A, D e G, são os rótulos dados pelo LLM.")
    for linha in pc.CABECALHO[1:]:
        _par(doc, linha, mono=True, size=10)
    _par(doc, "Início", mono=True, size=10)
    _tabela(doc, [(str(n), t[4:]) for n, t in pc.LINHAS],
            [1.0, 16.0], mono_cols=(1,))
    _par(doc, "Fim.", mono=True, size=10)

    _par(doc, "2. Valores usados nas execuções", bold=True, size=12)
    _par(doc, "Os valores vêm de PERCEPTRON_PARAMS, em src/protocolo.py, e dos valores padrão da classe. São os mesmos "
              "em todas as chamadas do código.")
    _tabela(doc, pc.PARAMETROS, [2.5, 3.5, 2.0, 9.0], header=("Símbolo", "Nome no código", "Valor", "Significado"))

    _par(doc, "3. Observações", bold=True, size=12)
    for obs in pc.OBSERVACOES:
        _par(doc, obs)

    _par(doc, "4. Correspondência com o código", bold=True, size=12)
    _par(doc, "Linhas do pseudocódigo e trechos correspondentes de src/relaxed_perceptron.py (código lido do arquivo "
              "no momento da geração). As linhas de fechamento (Fim_Se, Fim_Para, Fim_Enquanto) não estão na tabela.")
    _tabela(doc, [(str(n), ls, cod) for n, ls, cod in pc.tabela_correspondencia()], [1.3, 2.7, 13.0],
            mono_cols=(2,), header=("Linha", "Linhas do .py", "Código Python"))

    doc.save(SAIDA)
    print(f"Documento gerado: {SAIDA} ({os.path.getsize(SAIDA) / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
